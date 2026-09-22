"""Mahaller: mimari plandaki duvarlardan mahal sınırı, daire / mahal hiyerarşisi, mahal bazında keşif."""
from __future__ import annotations

import ezdxf
import pytest


def _daire_dxf(path):
    """Daire 1 (12×8 m) içinde SALON + HOL, yanında ORTAK KORİDOR; kameralar ve prizler yerleştirilmiş.

    HOL ile SALON arasındaki duvarda 1 m kapı boşluğu var: mahal ancak boşluk köprülenirse kapanır."""
    doc = ezdxf.new("R2010"); doc.header["$INSUNITS"] = 6          # metre
    for n in ("DUVAR", "YAZI", "KSF-ZAY-KAMERA-DOME", "KSF-ELK-PRIZ-TOPRAKLI"):
        doc.layers.add(n)
    doc.blocks.new(name="KAM").add_circle((0, 0), 0.15)
    doc.blocks.new(name="PRZ").add_circle((0, 0), 0.10)
    msp = doc.modelspace()
    msp.add_lwpolyline([(0, 0), (12, 0), (12, 8), (0, 8)], close=True, dxfattribs={"layer": "DUVAR"})
    msp.add_line((8, 0), (8, 3.5), dxfattribs={"layer": "DUVAR"})   # SALON | HOL ayıran duvar
    msp.add_line((8, 4.5), (8, 8), dxfattribs={"layer": "DUVAR"})   # (arada 1 m kapı boşluğu)
    msp.add_lwpolyline([(13, 0), (20, 0), (20, 8), (13, 8)], close=True, dxfattribs={"layer": "DUVAR"})
    for t, (x, y) in [("DAİRE 1", (3, 7)), ("SALON\\P64.00 m2", (3, 4)),
                      ("HOL\\P32.00 m2", (10, 4)), ("ORTAK KORİDOR\\P56.00 m2", (16, 4))]:
        msp.add_text(t, dxfattribs={"layer": "YAZI", "height": 0.3}).set_placement((x, y))
    for x, y in [(9, 2), (11, 6)]:                                  # HOL: 2 kamera
        msp.add_blockref("KAM", (x, y), dxfattribs={"layer": "KSF-ZAY-KAMERA-DOME"})
    for x, y in [(2, 2), (5, 6), (6, 2)]:                           # SALON: 3 kamera
        msp.add_blockref("KAM", (x, y), dxfattribs={"layer": "KSF-ZAY-KAMERA-DOME"})
    msp.add_blockref("KAM", (16, 2), dxfattribs={"layer": "KSF-ZAY-KAMERA-DOME"})   # koridor: 1
    for x, y in [(1, 1), (2, 7), (6, 7), (7, 1)]:                   # SALON: 4 priz
        msp.add_blockref("PRZ", (x, y), dxfattribs={"layer": "KSF-ELK-PRIZ-TOPRAKLI"})
    msp.add_blockref("PRZ", (9, 7), dxfattribs={"layer": "KSF-ELK-PRIZ-TOPRAKLI"})  # HOL: 1 priz
    doc.saveas(path)
    return path


def test_detect_spaces_hierarchy(tmp_path):
    """Duvarlardan mahal: kapı boşluğu köprülenir, daire içindeki odalar onun çocuğu olur."""
    from app.parser.loader import load_dxf
    from app.parser.spaces import detect_spaces
    d = load_dxf(str(_daire_dxf(tmp_path / "daire.dxf")))
    spaces, warns = detect_spaces(d, ["DUVAR"])
    by = {s.name: s for s in spaces}
    assert set(by) == {"DAİRE 1", "SALON", "HOL", "ORTAK KORİDOR"}
    assert by["DAİRE 1"].kind == "grup" and by["SALON"].kind == "mahal"
    assert by["SALON"].parent == by["DAİRE 1"].index and by["HOL"].parent == by["DAİRE 1"].index
    assert by["ORTAK KORİDOR"].parent is None                  # daire dışında, ayrı mahal
    # çizimden ölçülen alan ile yazıdaki alan birbirini doğruluyor
    assert by["SALON"].area == pytest.approx(64.0, abs=0.5) and by["SALON"].label_area == 64.0
    assert by["HOL"].area == pytest.approx(32.0, abs=0.5)      # kapı boşluğu köprülenmeseydi kapanmazdı
    assert by["DAİRE 1"].area == pytest.approx(96.0, abs=0.5)
    assert any("Mahal okundu" in w for w in warns)


def test_space_breakdown_api(client, tmp_path):
    """Mahal bazında keşif: kamera / priz noktası hangi mahaldeyse o mahale sayılır, daire toplanır."""
    p = _daire_dxf(tmp_path / "daire.dxf")
    pid = client.post("/api/projects", json={"name": "Mahal"}).json()["id"]
    with open(p, "rb") as f:
        r = client.post(f"/api/projects/{pid}/drawings", files={"file": ("ZEMİN KAT PLANI.dxf", f, "application/dxf")})
    assert r.status_code == 201, r.text
    out = client.get(f"/api/projects/{pid}/spaces").json()
    by = {s["name"]: s for s in out["spaces"]}
    assert set(by) >= {"DAİRE 1", "SALON", "HOL", "ORTAK KORİDOR"}
    assert by["HOL"]["path"] == "DAİRE 1 / HOL" and by["HOL"]["kind"] == "mahal"

    def qty(space, kind):
        return sum(i["quantity"] for i in space["items"] if i["kind"] == kind)

    assert qty(by["HOL"], "kamera") == 2 and qty(by["SALON"], "kamera") == 3
    assert qty(by["ORTAK KORİDOR"], "kamera") == 1
    assert qty(by["SALON"], "priz") == 4 and qty(by["HOL"], "priz") == 1
    # "evde 5 kamera": daire toplamı çocuklarının toplamıdır
    daire = {i["kind"]: i["quantity"] for i in by["DAİRE 1"]["total_items"]}
    assert daire["kamera"] == 5 and daire["priz"] == 5
    # koridor daireye dahil değil: proje toplamı 6 kamera
    assert sum(qty(s, "kamera") for s in out["spaces"] if s["kind"] == "mahal") == 6


def test_label_clustering_real_world(tmp_path):
    """Gerçek projede mahal etiketi PARÇALI yazılır: ad, kod ve alan ayrı TEXT nesneleridir.

    (Yat Kulübü uygulama projesinden alınan gerçek düzen: "RESTORAN" / "L_Z_01" / ":" / "309.49 m²")"""
    import ezdxf
    from app.parser.loader import load_dxf
    from app.parser.spaces import _labels
    doc = ezdxf.new("R2010"); doc.header["$INSUNITS"] = 6
    doc.layers.add("YAZI")
    msp = doc.modelspace()
    for i, (ad, kod, alan, y) in enumerate([("RESTORAN", "L_Z_01", "309.49 m²", 0.0),
                                            ("MERDİVEN", "L_Z_M01", "32.37 m²", 40.0),
                                            ("MERDİVEN", "L_Z_M02", "32.37 m²", 80.0)]):
        msp.add_text(ad, dxfattribs={"layer": "YAZI", "height": 0.24}).set_placement((0, y + 0.64))
        msp.add_text(kod, dxfattribs={"layer": "YAZI", "height": 0.24}).set_placement((-0.89, y))
        msp.add_text(":", dxfattribs={"layer": "YAZI", "height": 0.24}).set_placement((-0.4, y))
        msp.add_text(alan, dxfattribs={"layer": "YAZI", "height": 0.24}).set_placement((0, y))
    p = tmp_path / "etiket.dxf"; doc.saveas(p)
    labs = {l.code: l for l in _labels(load_dxf(str(p))) if l.area > 0}
    assert set(labs) == {"L_Z_01", "L_Z_M01", "L_Z_M02"}
    assert labs["L_Z_01"].name == "RESTORAN" and labs["L_Z_01"].area == 309.49
    # aynı adlı ve aynı alanlı iki merdiven AYRI mahaldir: kodları farklı
    assert labs["L_Z_M01"].name == labs["L_Z_M02"].name == "MERDİVEN"
    assert labs["L_Z_M01"].area == labs["L_Z_M02"].area == 32.37


def test_area_boundary_layer_offset(tmp_path):
    """Mahal alan sınırı ayrı katmanda ve planın AYRI BİR KOPYASINDA çizilmiş olabilir.

    (Yat Kulübü projesindeki gerçek durum: 'alan çizgisi' katmanındaki çokgenler, mahal
    yazılarının 240 m altındaki alan hesabı kopyasında duruyor.) Alan eşleşmesinden kayma
    bulunup çokgenler plana taşınır; alanlar yazıyı tutmazsa çokgen kullanılmaz."""
    import ezdxf
    from app.parser.loader import load_dxf
    from app.parser.spaces import detect_spaces
    doc = ezdxf.new("R2010"); doc.header["$INSUNITS"] = 6
    for n in ("YAZI", "alan çizgisi"):
        doc.layers.add(n)
    msp = doc.modelspace()
    DY = 240.0                                   # alan hesabı kopyasının kayması
    odalar = [("RESTORAN", "L_Z_01", 0.0, 0.0, 20.0, 15.0),      # 300 m²
              ("HOL", "L_Z_06", 22.0, 0.0, 10.0, 5.0),           # 50 m²
              ("MUTFAK", "L_Z_12", 34.0, 0.0, 8.0, 5.0)]         # 40 m²
    for ad, kod, x, y, w, h in odalar:
        msp.add_lwpolyline([(x, y - DY), (x + w, y - DY), (x + w, y + h - DY), (x, y + h - DY)],
                           close=True, dxfattribs={"layer": "alan çizgisi"})
        cx, cy = x + w / 2, y + h / 2            # etiket odanın ortasında, ASIL planda
        msp.add_text(ad, dxfattribs={"layer": "YAZI", "height": 0.24}).set_placement((cx, cy + 0.6))
        msp.add_text(kod, dxfattribs={"layer": "YAZI", "height": 0.24}).set_placement((cx - 0.9, cy))
        msp.add_text(f"{w * h:.2f} m²", dxfattribs={"layer": "YAZI", "height": 0.24}).set_placement((cx, cy))
    p = tmp_path / "alan.dxf"; doc.saveas(p)
    spaces, warns = detect_spaces(load_dxf(str(p)), ["alan çizgisi"])
    by = {s.code: s for s in spaces}
    assert set(by) == {"L_Z_01", "L_Z_06", "L_Z_12"}
    assert by["L_Z_01"].name == "RESTORAN" and by["L_Z_01"].area == pytest.approx(300.0, abs=0.5)
    # sınır çizimden ölçüldü ve yazıdaki alanı doğruluyor
    assert all(s.area_source == "drawing" for s in by.values())
    assert any("alan çizgilerinden alındı" in w and "240" in w for w in warns)


def test_mahal_name_filter():
    """Çizim işareti mahal adı sanılmasın: iki harfli kısaltma (DK), ölçü notu, kot."""
    from app.parser.spaces import _is_name
    assert _is_name("MUTFAK") and _is_name("ÇALIŞMA ODASI") and _is_name("WC") and _is_name("HOL")
    assert not _is_name("DK") and not _is_name("TK")            # çizim kısaltması
    assert not _is_name("30X(31 / 16.33)") and not _is_name("27X34")   # ölçü notu
    assert not _is_name("+4.15") and not _is_name("S1") and not _is_name("1/100")


def test_space_derived_items(client, tmp_path):
    """Mahal bazında türetme: şap / kaplama / tavan mahal alanından, sıva-boya ÖLÇÜLEN çevreden.

    HOL 4×8 m: gerçek çevre 24 m. Proje genelindeki "kare mahal" varsayımı (4·√alan) 22,6 m derdi;
    mahal sınırı bilindiğinde varsayıma gerek yok."""
    p = _daire_dxf(tmp_path / "daire.dxf")
    pid = client.post("/api/projects", json={"name": "Türetme", "storey_height": 3.0}).json()["id"]
    with open(p, "rb") as f:
        assert client.post(f"/api/projects/{pid}/drawings",
                           files={"file": ("ZEMİN KAT PLANI.dxf", f, "application/dxf")}).status_code == 201
    by = {s["name"]: s for s in client.get(f"/api/projects/{pid}/spaces").json()["spaces"]}
    hol = by["HOL"]
    assert hol["area_source"] == "drawing" and hol["perimeter"] == pytest.approx(24.0, abs=0.1)
    d = {i["kind"]: i for i in hol["derived"]}
    assert d["sap"]["quantity"] == pytest.approx(32.0 * 0.05)          # 5 cm varsayılan, notu yok
    assert "VARSAYILAN" in d["sap"]["note"]
    assert d["doseme_kaplama"]["quantity"] == pytest.approx(32.0)
    assert d["tavan_siva_boya"]["quantity"] == pytest.approx(32.0)
    # duvar yüzeyi: çevre × (kat yüksekliği − döşeme) — çevre ölçüldü, varsayılmadı
    assert d["siva"]["quantity"] == pytest.approx(24.0 * 2.85, rel=1e-3)
    assert d["boya"]["quantity"] == d["siva"]["quantity"]
    assert "ÖLÇÜLDÜ" in d["siva"]["note"]
    # bağımsız bölüm satırı çocuklarının türetilmiş kalemlerini de toplar
    daire = {i["kind"]: i["quantity"] for i in by["DAİRE 1"]["total_items"]}
    assert daire["tavan_siva_boya"] == pytest.approx(64.0 + 32.0)


def _mimari_dxf(path, dx=0.0, dy=0.0, tesisat=None):
    """Mimari plan (duvar + mahal yazısı) ya da aynı planın tesisat kopyası.

    tesisat: [(katman, [(x, y), …])] — elektrik / mekanik sembolleri. dx, dy ile pafta kaydırılır."""
    import ezdxf
    doc = ezdxf.new("R2010"); doc.header["$INSUNITS"] = 6
    for n in ["DUVAR", "YAZI"] + [k for k, _ in (tesisat or [])]:
        doc.layers.add(n)
    for blk in ("SEM",):
        if blk not in doc.blocks:
            doc.blocks.new(name=blk).add_circle((0, 0), 0.15)
    msp = doc.modelspace()
    M = lambda x, y: (x + dx, y + dy)                                    # noqa: E731
    msp.add_lwpolyline([M(0, 0), M(12, 0), M(12, 8), M(0, 8)], close=True, dxfattribs={"layer": "DUVAR"})
    msp.add_line(M(8, 0), M(8, 3.5), dxfattribs={"layer": "DUVAR"})
    msp.add_line(M(8, 4.5), M(8, 8), dxfattribs={"layer": "DUVAR"})
    for ad, alan, x, y in [("SALON", 64.0, 4, 4), ("HOL", 32.0, 10, 4)]:
        msp.add_text(ad, dxfattribs={"layer": "YAZI", "height": 0.24}).set_placement(M(x, y + 0.6))
        msp.add_text(f"{alan:.2f} m²", dxfattribs={"layer": "YAZI", "height": 0.24}).set_placement(M(x, y))
    for i in range(22):
        msp.add_text(f"N{i}", dxfattribs={"layer": "YAZI", "height": 0.2}).set_placement(M(0.5 + i * 0.5, 9))
    for katman, noktalar in (tesisat or []):
        for x, y in noktalar:
            msp.add_blockref("SEM", M(x, y), dxfattribs={"layer": katman})
    doc.saveas(path)
    return path


@pytest.mark.parametrize("kayma", [(0.0, 0.0), (250.0, -120.0)])
def test_spaces_across_sheets(client, tmp_path, kayma):
    """Mahal mimariden çıkar; AYRI paftadaki elektrik ve mekanik kalemleri o mahallere yazılır.

    Tesisat paftası başka koordinatta çizilmiş olsa bile mahal yazılarından kayma bulunur ve
    doğrulanır (o kaymayla kaç eleman mahalin içine düşüyor)."""
    dx, dy = kayma
    mim = _mimari_dxf(tmp_path / "mimari.dxf")
    elk = _mimari_dxf(tmp_path / "elektrik.dxf", dx, dy, tesisat=[
        ("KSF-ELK-ARMATUR-LED", [(2, 2), (5, 6), (6, 2), (3, 5)]),      # SALON: 4 armatür
        ("KSF-ELK-PRIZ-TOPRAKLI", [(1, 1), (7, 1)]),                    # SALON: 2 priz
        ("KSF-ZAY-KAMERA-DOME", [(9, 2), (11, 6)]),                     # HOL: 2 kamera
    ])
    mek = _mimari_dxf(tmp_path / "mekanik.dxf", dx, dy, tesisat=[
        ("KSF-MEK-VRF_IC_UNITE-5.6KW", [(4, 7)]),                       # SALON: 1 iç ünite
        ("KSF-MEK-MENFEZ-ANEMOSTAT", [(3, 3), (6, 5), (10, 3)]),        # SALON 2, HOL 1
    ])
    pid = client.post("/api/projects", json={"name": "Disiplinler", "storey_height": 3.0}).json()["id"]
    for ad, yol in [("ZEMİN KAT PLANI.dxf", mim), ("ZEMİN KAT ELEKTRİK PLANI.dxf", elk),
                    ("ZEMİN KAT MEKANİK PLANI.dxf", mek)]:
        with open(yol, "rb") as f:
            assert client.post(f"/api/projects/{pid}/drawings",
                               files={"file": (ad, f, "application/dxf")}).status_code == 201, ad
    out = client.get(f"/api/projects/{pid}/spaces").json()
    by = {s["name"]: s for s in out["spaces"]}
    assert {"SALON", "HOL"} <= set(by)

    def adet(mahal, kind):
        return sum(i["quantity"] for i in by[mahal]["items"] if i["kind"] == kind)

    assert adet("SALON", "armatur") == 4 and adet("SALON", "priz") == 2
    assert adet("HOL", "kamera") == 2 and adet("SALON", "kamera") == 0
    assert adet("SALON", "vrf_ic_unite") == 1
    assert adet("SALON", "menfez") == 2 and adet("HOL", "menfez") == 1
    # hizalama raporlanır: hangi pafta hangi mahal setine, ne kadar kaymayla
    hiz = {h["drawing"]: h for h in out["alignment"]}
    assert len(hiz) == 2
    for h in hiz.values():
        assert h["hit"] == h["total"] and (abs(h["dx"] + dx) < 0.05 or dx == 0)


def test_spaces_xlsx(client, tmp_path):
    """Mahal metrajı Excel'e dökülür: mahal listesi, mahal × kalem, mahale girmeyenler, hizalama."""
    from openpyxl import load_workbook
    from io import BytesIO
    mim = _mimari_dxf(tmp_path / "mimari.dxf")
    elk = _mimari_dxf(tmp_path / "elk.dxf", 0, 0, tesisat=[
        ("KSF-ELK-ARMATUR-LED", [(2, 2), (5, 6), (6, 2)]), ("KSF-ZAY-KAMERA-DOME", [(9, 2)])])
    pid = client.post("/api/projects", json={"name": "Excel", "storey_height": 3.0}).json()["id"]
    for ad, yol in [("ZEMİN KAT PLANI.dxf", mim), ("ZEMİN KAT AYDINLATMA PLANI.dxf", elk)]:
        with open(yol, "rb") as f:
            client.post(f"/api/projects/{pid}/drawings", files={"file": (ad, f, "application/dxf")})
    r = client.get(f"/api/projects/{pid}/spaces.xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    wb = load_workbook(BytesIO(r.content))
    assert {"Mahaller", "Mahal metrajı", "Mahale girmeyen"} <= set(wb.sheetnames)
    mahaller = [[c.value for c in row] for row in wb["Mahaller"].iter_rows(min_row=4)]
    assert {r[1] for r in mahaller} == {"SALON", "HOL"}
    assert all(r[5] == "çizimden ölçüldü" for r in mahaller)          # sınır doğrulandı
    metraj = [[c.value for c in row] for row in wb["Mahal metrajı"].iter_rows(min_row=4)]
    armatur = [r for r in metraj if r[1] == "SALON" and "armat" in (r[5] or "").lower()]
    assert armatur and armatur[0][7] == 3                              # SALON: 3 armatür
    assert any(r[1] == "HOL" and "Kamera" in (r[5] or "") and r[7] == 1 for r in metraj)
    assert any("Şap" in (r[5] or "") and r[8] == "türetildi" for r in metraj)


def test_alignment_ignores_lines(client, tmp_path):
    """Hizalama yalnız NOKTA elemanlarıyla aranır: kablo / boru gibi çizgiler mahaller arasında
    uzandığı için aramayı yanıltır — varlıkları bulunan kaymayı değiştirmemeli."""
    import ezdxf
    from app.services import align_drawing, _space_polys
    from app.parser.loader import load_dxf
    from app.parser.spaces import detect_spaces
    from types import SimpleNamespace as NS
    mim = load_dxf(str(_mimari_dxf(tmp_path / "m.dxf")))
    sps, _ = detect_spaces(mim, ["DUVAR"])
    polys = _space_polys([s.to_dict() for s in sps])
    DX, DY = 640.0, -275.0                                    # tesisat paftası başka koordinatta
    nokta = [NS(points=[[x + DX, y + DY]], etype="fixture", length=0.0, area=0.0, count=1)
             for x, y in [(2, 2), (5, 6), (6, 2), (3, 5), (1, 1), (7, 1), (9, 2), (11, 6), (10, 5)]]
    # mahalleri boydan boya geçen hatlar: nokta değildir, hizalamaya girmemeli
    hat = [NS(points=[[0 + DX, y + DY], [12 + DX, y + DY]], etype="cable", length=12.0, area=0.0, count=1)
           for y in (1.0, 2.0, 3.0, 5.0, 6.0, 7.0)]
    yalniz = align_drawing(NS(rooms=[]), NS(rooms=[]), nokta, polys)
    karisik = align_drawing(NS(rooms=[]), NS(rooms=[]), nokta + hat, polys)
    # kayma, noktaları doğru mahallere oturtacak kadar doğru bulunur (birebir tek bir değer değildir:
    # noktalar mahalin içinde kaldığı sürece küçük oynamalar aynı sonucu verir)
    from shapely.geometry import Point as SPoint
    yer = []
    for e in nokta:
        p0 = SPoint(e.points[0][0] + yalniz["dx"], e.points[0][1] + yalniz["dy"])
        yer.append(next((sp["name"] for sp, g in polys if g.contains(p0)), None))
    assert yer == ["SALON"] * 6 + ["HOL"] * 3
    assert (karisik["dx"], karisik["dy"]) == (yalniz["dx"], yalniz["dy"])   # hatlar kaymayı değiştirmedi
    assert karisik["total"] == len(nokta) == yalniz["hit"]                  # ve sayıma girmedi
