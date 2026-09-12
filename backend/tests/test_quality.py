from datetime import datetime
from io import BytesIO
from types import SimpleNamespace as Obj

from openpyxl import load_workbook
from app.quality import build_quality


def drawing(**kw):
    defaults = dict(id=1, label="Zemin", filename="zemin.dxf", plan_type="sta_kat_kalip",
                    analyzed_at=datetime.now(), unit_detected=True, unit_override=None, warnings=[])
    return Obj(**(defaults | kw))


def test_empty_and_unresolved_geometry_are_incomplete():
    q = build_quality([drawing(warnings=["2 kiriş etiketi hiçbir kirişe atanmadı"])], [], [], {}, {}, {})
    assert q['status'] == 'incomplete'
    assert {'empty_drawing', 'unresolved_geometry'} <= {i['code'] for i in q['issues']}
    assert not q['certified']


def test_non_measured_sheets_not_reported_as_empty():
    q = build_quality([drawing(plan_type='mim_kesit')], [], [], {}, {}, {})
    assert 'empty_drawing' not in {i['code'] for i in q['issues']}
    assert q['status'] == 'review_required' and not q['certified']


def test_assumptions_and_missing_labor_are_visible():
    e = Obj(drawing_id=1, included=True, confidence=1, warnings=[])
    it = Obj(kind='beton', detail={'recipe': True})
    q = build_quality([drawing()], [e], [it], {'groups': [{'rebar_ratio_kg': 2500}]},
                      {'concrete_class': 'C40/50'}, {}, {'missing_labor': ['beton:*']})
    assumptions = {a['key']: a for a in q['assumptions']}
    assert assumptions['concrete_class']['source'] == 'user'
    assert assumptions['concrete_waste_pct']['source'] == 'default'
    assert q['estimated_rebar_kg'] == 2500
    assert q['derived_counts']['recipe'] == 1
    assert q['status'] == 'incomplete'
    assert 'missing_labor' in {i['code'] for i in q['issues']}


def test_quality_api_and_export(client, storey_dxf):
    pid = client.post('/api/projects', json={'name': 'Kontrol'}).json()['id']
    empty = client.get(f'/api/projects/{pid}/quantities').json()['quality']
    assert empty['status'] == 'incomplete'
    with open(storey_dxf, 'rb') as f:
        assert client.post(f'/api/projects/{pid}/drawings', files={'file': ('kat_plani.dxf', f)}).status_code == 201
    q = client.get(f'/api/projects/{pid}/quantities').json()['quality']
    assert q['estimated_rebar_kg'] > 0
    cost = client.get(f'/api/projects/{pid}/cost').json()
    assert cost['quality']['status'] == 'incomplete'
    response = client.get(f'/api/projects/{pid}/cost.xlsx')
    assert response.status_code == 200
    wb = load_workbook(BytesIO(response.content))
    assert wb.sheetnames[0] == 'Kontrol'
    assert wb['Kontrol']['A1'].value == 'HESAP TASLAĞI'
    assert 'HESAP TASLAĞI' in wb['Keşif']['A3'].value

