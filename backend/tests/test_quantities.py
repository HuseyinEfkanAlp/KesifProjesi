import pytest
from openpyxl import load_workbook
from io import BytesIO

from app.cost.pricing import PriceItem, compute_cost, default_price_items
from app.export.excel import build_workbook
from app.parser.analyzer import analyze_file
from app.quantity.boq import structural_items
from app.quantity.engine import ElementData, QuantityParams, compute_all, compute_element
from app.quantity.summary import summarize


def test_column_formulas():
    p = QuantityParams(storey_height=3.0, slab_thickness=0.15, storey_count=2)
    el = ElementData(id=1, etype="column", area=0.18, perimeter=1.8)
    ln = compute_element(el, p)
    assert ln.concrete_m3 == pytest.approx(0.18 * 2.85)
    assert ln.formwork_m2 == pytest.approx(1.8 * 2.85)
    assert ln.rebar_kg == pytest.approx(0.18 * 2.85 * 180)
    assert ln.multiplier == 2
    assert ln.total_concrete == pytest.approx(0.18 * 2.85 * 2)


def test_beam_slab_wall_foundation_formulas():
    p = QuantityParams(storey_height=3.0, slab_thickness=0.15, storey_count=3)
    beam = compute_element(ElementData(id=1, etype="beam", b=0.25, h=0.50, length=4.7), p)
    assert beam.concrete_m3 == pytest.approx(0.25 * 0.35 * 4.7)
    assert beam.formwork_m2 == pytest.approx(2 * 0.35 * 4.7)          # brüt döşeme: alt yüz döşeme kalıbında, yalnız yanlar (h − d)
    pn = QuantityParams(storey_height=3.0, slab_thickness=0.15, storey_count=3, beam_full_height=True)
    beam_n = compute_element(ElementData(id=1, etype="beam", b=0.25, h=0.50, length=4.7), pn)
    assert beam_n.concrete_m3 == pytest.approx(0.25 * 0.50 * 4.7)
    assert beam_n.formwork_m2 == pytest.approx((0.25 + 2 * 0.35) * 4.7)   # net döşeme: alt yüz + yanlar döşeme altına kadar
    slab = compute_element(ElementData(id=2, etype="slab", area=34.98, thickness=0.15), p)
    assert slab.concrete_m3 == pytest.approx(34.98 * 0.15)
    assert slab.formwork_m2 == pytest.approx(34.98)
    wall = compute_element(ElementData(id=3, etype="shear_wall", area=0.5, length=2.5), p)
    assert wall.concrete_m3 == pytest.approx(0.5 * 2.85)
    assert wall.formwork_m2 == pytest.approx(2 * 2.5 * 2.85)
    raft = compute_element(ElementData(id=4, etype="foundation", subtype="raft", area=42.0, perimeter=26.0, thickness=0.6), p)
    assert raft.concrete_m3 == pytest.approx(25.2)
    assert raft.formwork_m2 == pytest.approx(26.0 * 0.6)
    assert raft.multiplier == 1  # temel kat sayısıyla çarpılmaz
    strip = compute_element(ElementData(id=5, etype="foundation", subtype="strip", b=0.6, h=0.8, length=4.0), p)
    assert strip.concrete_m3 == pytest.approx(0.6 * 0.8 * 4.0)
    assert strip.formwork_m2 == pytest.approx(2 * 4.0 * 0.8)
    missing = compute_element(ElementData(id=6, etype="beam", b=0.25, h=None, length=4.7), p)
    assert missing.concrete_m3 == 0 and missing.notes
    # etiketsiz kiriş, kattaki baskın kiriş yüksekliği biliniyorsa onunla hesaplanır
    p2 = QuantityParams(storey_height=3.0, slab_thickness=0.15, beam_depth=0.45, beam_full_height=True)
    fb = compute_element(ElementData(id=7, etype="beam", b=0.25, h=None, length=4.0), p2)
    assert fb.concrete_m3 == pytest.approx(0.25 * 0.45 * 4.0) and "baskın" in fb.notes[0]


def test_end_to_end_storey(storey_dxf):
    r = analyze_file(str(storey_dxf))
    p = QuantityParams(storey_height=3.0, slab_thickness=0.15, storey_count=1)
    lines = compute_all([ElementData.from_obj(e, id=i) for i, e in enumerate(r.elements)], p)
    s = summarize(lines)
    groups = {g["key"]: g for g in s["groups"]}
    assert groups["column"]["concrete_m3"] == pytest.approx(4 * 0.18 * 2.85, abs=1e-3)   # 2.052
    assert groups["beam"]["concrete_m3"] == pytest.approx(0.25 * 0.35 * (4.7 + 5.4), abs=1e-3)
    assert groups["slab"]["concrete_m3"] == pytest.approx(34.98 * 0.15, abs=1e-3)
    assert groups["shear_wall"]["concrete_m3"] == pytest.approx(0.5 * 2.85, abs=1e-3)
    expected_total = 4 * 0.18 * 2.85 + 0.25 * 0.35 * 10.1 + 34.98 * 0.15 + 0.5 * 2.85
    assert s["totals"]["concrete_m3"] == pytest.approx(expected_total, abs=1e-2)

    items = structural_items(s)
    prices = default_price_items(items)
    assert any(pi.key == "beton:column" for pi in prices)
    cost = compute_cost(items, [PriceItem("beton:*", "Beton", "m³", 4000), PriceItem("beton:column", "Beton kolon", "m³", 4500),
                                PriceItem("demir:*", "Demir", "kg", 30)], vat_rate=0.2)
    col_line = next(l for l in cost["lines"] if l["key"] == "beton:column")
    assert col_line["unit_price"] == 4500 and col_line["price_source"] == "özel"
    slab_line = next(l for l in cost["lines"] if l["key"] == "beton:slab")
    assert slab_line["unit_price"] == 4000 and slab_line["price_source"] == "genel"
    assert all(l["unit_price"] == 0 for l in cost["lines"] if l["kind"] == "kalip")
    assert cost["grand_total"] == pytest.approx(cost["subtotal"] * 1.2, abs=0.05)

    xlsx = build_workbook({"name": "Test", "storey_height": 3.0, "slab_thickness": 0.15}, lines, s, cost,
                          boq=[i.to_dict() for i in items])
    wb = load_workbook(BytesIO(xlsx))
    assert wb.sheetnames == ["Keşif", "Statik Özet", "Kat Bazında", "Eleman Metrajı", "Maliyet"]
    assert wb["Eleman Metrajı"].max_row == 1 + len(lines)
    assert wb["Keşif"].max_row >= 4 + len(items) + 3      # kalemler + tür toplamları bloğu
    assert any(c.value == "TÜR TOPLAMLARI" for row in wb["Keşif"].iter_rows(min_col=1, max_col=1) for c in row)


def test_summary_sections_group_same_section():
    """Aynı kesitteki elemanlar kesit bazında tek satırda toplanır."""
    from app.quantity.engine import ElementData, QuantityParams, compute_all
    from app.quantity.summary import summarize
    els = [ElementData(id=i, etype="beam", name=f"K{i}", subtype=None, b=0.3, h=0.6, thickness=None, area=0.3 * 4.0, length=4.0,
                       perimeter=0, count=1) for i in range(5)]
    els += [ElementData(id=10 + i, etype="column", name=f"S{i}", subtype=None, b=0.4, h=0.4, thickness=None, area=0.16, length=0,
                        perimeter=1.6, count=1) for i in range(3)]
    lines = compute_all(els, QuantityParams(storey_height=3.0, slab_thickness=0.15, storey_count=2))
    info = {e.id: {"b": e.b, "h": e.h, "thickness": e.thickness, "area": e.area, "length": e.length} for e in els}
    s = summarize(lines, [], info)
    secs = {x["key"]: x for x in s["sections"]}
    assert secs["beam|30x60"]["element_count"] == 10 and secs["beam|30x60"]["length_m"] == pytest.approx(40.0)
    col_group = next(g for g in s["groups"] if g["etype"] == "column")
    assert secs["column|40x40"]["element_count"] == 6 and secs["column|40x40"]["concrete_m3"] == pytest.approx(col_group["concrete_m3"], abs=0.01)
    assert len(s["sections"]) == 2
