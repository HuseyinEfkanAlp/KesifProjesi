"""Mahal metrajı tablosu (Excel): her mahalin kalemleri, disiplin gruplarıyla.

Keşif tablosu "bina toplamı"dır; bu tablo **mahal bazında** döker: hangi mahalde ne kadar seramik,
şap, armatür, priz, kamera, iç ünite, menfez var. Kaynak `services.space_breakdown`.
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

BASLIK = PatternFill("solid", fgColor="E8EEF7")


def _basliklar(ws, satir: int, kolonlar: list[str]) -> None:
    for i, ad in enumerate(kolonlar, start=1):
        h = ws.cell(row=satir, column=i, value=ad)
        h.font = Font(bold=True)
        h.fill = BASLIK
        h.alignment = Alignment(vertical="center", wrap_text=True)


def build_spaces_workbook(project: dict, breakdown: dict) -> bytes:
    """Mahal metrajı çalışma kitabı: (1) mahal listesi, (2) mahal × kalem dökümü, (3) atanamayanlar."""
    wb = Workbook()
    spaces = breakdown.get("spaces") or []

    ws = wb.active
    ws.title = "Mahaller"
    ws["A1"] = f"Proje: {project.get('name', '')} — mahal listesi"
    ws["A1"].font = Font(bold=True, size=13)
    _basliklar(ws, 3, ["Mahal kodu", "Mahal", "Bağlı olduğu bölüm", "Alan (m²)", "Çevre (m)",
                       "Alan kaynağı", "Yazıdaki alan (m²)", "Pafta"])
    for sp in sorted(spaces, key=lambda s: (s.get("drawing") or "", -float(s.get("area") or 0))):
        kaynak = {"drawing": "çizimden ölçüldü", "label": "yalnız mahal yazısından",
                  "polygon": "çizimden (yazıda alan yok)"}.get(sp.get("area_source"), sp.get("area_source") or "")
        ws.append([sp.get("code") or "", sp.get("name") or "", sp.get("path") or "", round(float(sp.get("area") or 0), 2),
                   round(float(sp.get("perimeter") or 0), 2) or "", kaynak,
                   round(float(sp.get("label_area") or 0), 2) or "", sp.get("drawing") or ""])
    for c, w in zip("ABCDEFGH", (14, 26, 30, 12, 11, 24, 16, 28)):
        ws.column_dimensions[c].width = w

    ws2 = wb.create_sheet("Mahal metrajı")
    ws2["A1"] = f"Proje: {project.get('name', '')} — mahal bazında keşif"
    ws2["A1"].font = Font(bold=True, size=13)
    _basliklar(ws2, 3, ["Kat / Pafta", "Mahal kodu", "Mahal", "Alan (m²)", "İş grubu", "Disiplin", "Kalem",
                        "Birim", "Miktar", "Kaynak", "Not"])
    for sp in sorted(spaces, key=lambda s: (s.get("drawing") or "", -float(s.get("area") or 0))):
        for it in (sp.get("items") or []) + (sp.get("derived") or []):
            ws2.append([sp.get("drawing") or "", sp.get("code") or "", sp.get("name") or "",
                        round(float(sp.get("area") or 0), 2),
                        it.get("work_group_label") or "", it.get("discipline_label") or "", it.get("label") or "",
                        it.get("unit") or "", it.get("quantity") or 0,
                        "türetildi" if it.get("derived") else "çizimden ölçüldü",
                        it.get("note") or "; ".join(it.get("notes") or [])])
    for c, w in zip("ABCDEFGHIJK", (30, 14, 24, 11, 12, 20, 34, 8, 12, 18, 60)):
        ws2.column_dimensions[c].width = w

    ws3 = wb.create_sheet("Mahale girmeyen")
    ws3["A1"] = breakdown.get("unassigned_reason") or "Mahale atanamayan kalemler"
    ws3["A1"].font = Font(bold=True)
    # Kapsam notu: bir kalemin burada olmaması eksiklik değil, o kalemin doğası olabilir.
    if breakdown.get("scope_note"):
        ws3["A2"] = str(breakdown["scope_note"]).replace("**", "")
    _basliklar(ws3, 3, ["İş grubu", "Disiplin", "Kalem", "Birim", "Miktar"])
    for it in (breakdown.get("unassigned") or []):
        ws3.append([it.get("work_group_label") or "", it.get("discipline_label") or "", it.get("label") or "",
                    it.get("unit") or "", it.get("quantity") or 0])
    for c, w in zip("ABCDE", (12, 20, 34, 8, 12)):
        ws3.column_dimensions[c].width = w

    if breakdown.get("alignment"):
        ws4 = wb.create_sheet("Pafta hizalama")
        _basliklar(ws4, 1, ["Pafta", "Mahalleri veren pafta", "Kayma X (m)", "Kayma Y (m)",
                            "Nasıl bulundu", "Mahale düşen", "Toplam eleman"])
        for a in breakdown["alignment"]:
            ws4.append([a.get("drawing"), a.get("to"), a.get("dx"), a.get("dy"), a.get("how"),
                        a.get("hit"), a.get("total")])
        for c, w in zip("ABCDEFG", (34, 30, 13, 13, 26, 13, 14)):
            ws4.column_dimensions[c].width = w

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
