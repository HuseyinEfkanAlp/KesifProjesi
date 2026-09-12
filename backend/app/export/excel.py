"""Keşif, metraj ve maliyet Excel raporu (openpyxl)."""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..parser.layer_profile import ALL_ELEMENT_TYPES
from ..quantity.engine import QuantityLine
from ..quantity.summary import SUBTYPE_LABELS

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
MONEY = "#,##0.00"


def _header(ws, row: int, headers: list[str]) -> None:
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autosize(ws) -> None:
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 10), 45)


def build_workbook(project: dict, lines: list[QuantityLine], summary: dict, cost: dict,
                   element_info: dict | None = None, boq: list[dict] | None = None, quality: dict | None = None,
                   rebar: dict | None = None, headline: list | None = None) -> bytes:
    """element_info: element_id -> {"drawing": str, "layer": str, "b":..., "h":..., ...} (isteğe bağlı).
    boq: keşif kalemleri (tüm disiplinler). rebar: çap bazında demir raporu (quantity/rebar_report).
    headline: ana kalemler (quantity/headline)."""
    element_info = element_info or {}
    boq = boq or []
    wb = Workbook()

    # ---- Keşif Özeti (tüm disiplinler) ----
    ws = wb.active
    ws.title = "Keşif"
    ws["A1"] = f"Proje: {project.get('name', '')}"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = (f"Kat yüksekliği: {project.get('storey_height', '')} m   Döşeme kalınlığı: {project.get('slab_thickness', '')} m"
                f"   Duvar yüksekliği: {project.get('wall_height') or 'H − d'}   Günlük çalışma: {project.get('work_hours_per_day', 8)} saat")
    _header(ws, 4, ["İş grubu", "Poz", "Disiplin", "Tür", "Kalem", "Birim", "Miktar", "Adet / hat", "Not"])
    for it in boq:
        ws.append([it.get("work_group_label") or "", it.get("poz") or "", it["discipline_label"], it["kind_label"], it["label"],
                   it["unit"], it["quantity"], it.get("count") or "", "; ".join(it.get("notes") or [])])
    # tür toplamları (duvar, cam, kapı, körkasa…): sistem başlığı ve bilgi satırları hariç
    totals: dict[str, list] = {}
    for it in boq:
        d = it.get("detail") or {}
        if d.get("system") or d.get("info"):
            continue
        t = totals.setdefault(it["kind"], [it.get("work_group_label") or "", it["kind_label"], it["unit"], 0.0, 0])
        t[3] += float(it["quantity"] or 0)
        t[4] += 1
    if totals:
        ws.append([])
        ws.append(["TÜR TOPLAMLARI"])
        ws.append(["İş grubu", "Tür", "Birim", "Miktar", "Kalem sayısı"])
        for t in totals.values():
            ws.append([t[0], t[1], t[2], round(t[3], 3), t[4]])
    _autosize(ws)

    if quality:
        ws["A3"] = "HESAP TASLAĞI — " + quality["label"] + "; Kontrol sayfasını inceleyin."
        ws["A3"].font = Font(bold=True, color="9C2B1B")
        review = wb.create_sheet("Kontrol", 0)
        review.append(["HESAP TASLAĞI", quality["label"]])
        review.append([quality["notice"]])
        _header(review, 4, ["Önem", "Pafta", "Kontrol"])
        for issue in quality["issues"]:
            review.append(["Eksik" if issue["severity"] == "blocking" else "İncelenmeli",
                           issue["drawing"] or "Proje", issue["message"]])
        review.append([])
        review.append(["PARAMETRE", "DEĞER", "KAYNAK"])
        for assumption in quality["assumptions"]:
            review.append([assumption["label"], assumption["value"],
                           {"user": "Kullanıcı girişi", "drawing": "Kotlardan türetildi"}.get(assumption["source"], "Program varsayılanı")])
        _autosize(review)
        review.column_dimensions["C"].width = 100
        for row in review:
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        review.freeze_panes = "A5"

    # ---- Statik Metraj Özeti ----
    if summary.get("groups"):
        ws1 = wb.create_sheet("Statik Özet")
        _header(ws1, 1, ["Eleman Grubu", "Adet", "Beton (m³)", "Kalıp (m²)", "Demir (kg)"])
        for g in summary["groups"]:
            ws1.append([g["label"], g["element_count"], g["concrete_m3"], g["formwork_m2"], g["rebar_kg"]])
        t = summary["totals"]
        ws1.append(["TOPLAM", "", t["concrete_m3"], t["formwork_m2"], t["rebar_kg"]])
        for c in ws1[ws1.max_row]:
            c.font = BOLD
        _autosize(ws1)

    # ---- Ana kalemler (beton / kalıp / demir / duvar / sıva / boya) ----
    # "Ne kadar, neyden, ne kadarı ölçüldü" — keşif listesinin 70+ satırı arasında kaybolmasın.
    if headline:
        wsh = wb.create_sheet("Ana Kalemler", 1)
        wsh["A1"] = "Ana kalemler"
        wsh["A1"].font = Font(bold=True, size=13)
        _header(wsh, 3, ["Kalem", "Toplam", "Birim", "Ölçülen", "Oranla tahmin", "Fire",
                         "Brüt (duvar)", "Düşülen boşluk", "Çizimde kesilmiş", "İşçilik (saat)"])
        for h in headline:
            wsh.append([h["label"], round(h["total"], 2), h["unit"], round(h["net"], 2),
                        round(h.get("estimated") or 0, 2) or None, round(h.get("waste") or 0, 2) or None,
                        round(h["gross"], 2) if h.get("gross") is not None else None,
                        round(h["deducted"], 2) if h.get("deducted") is not None else None,
                        round(h["already_net"], 2) if h.get("already_net") else None,
                        round(h.get("labour_hours") or 0, 1) or None])
        for h in headline:
            wsh.append([])
            wsh.append([f"{h['label']} dökümü"])
            wsh[wsh.max_row][0].font = BOLD
            if h["kind"] == "duvar":
                _header(wsh, wsh.max_row + 1, ["Malzeme / kalınlık", f"Net ({h['unit']})", "Brüt", "Düşülen boşluk", "Çizimde kesilmiş"])
                for r in h["rows"]:
                    wsh.append([r["label"], round(r["quantity"], 2), r.get("gross_m2"),
                                r.get("openings_m2") or None, r.get("already_net_m2") or None])
            else:
                _header(wsh, wsh.max_row + 1, ["Kalem", f"Miktar ({h['unit']})", "Adet"])
                for r in h["rows"]:
                    wsh.append([r["label"], round(r["quantity"], 2), round(r.get("count") or 0) or None])
            for l in h.get("labour", []):
                wsh.append([f"  işçilik: {l['label']}", round(l["hours"], 1), "saat"])
            for e in h.get("extras", []):
                wsh.append([f"  sarf: {e['label']}", round(e["quantity"], 2), e["unit"]])
        _autosize(wsh)

    # ---- Demir (çap bazında sipariş + işçilik) ----
    # Sahada en çok sorulan sayı: "toplam kaç ton, hangi çaptan kaç kilo". Keşif sayfasında satırlar arasında
    # kaybolmasın diye kendi sayfası var.
    if rebar and rebar.get("rows"):
        wsr = wb.create_sheet("Demir")
        wsr["A1"] = "Demir siparişi ve işçiliği"
        wsr["A1"].font = Font(bold=True, size=13)
        tr = rebar["totals"]
        wsr["A2"] = (f"Sipariş (fire dahil): {tr['order_kg'] / 1000:,.1f} ton   "
                     f"Metraj: {tr['metraj_kg'] / 1000:,.1f} t   Oranla tahmin: {tr['ratio_kg'] / 1000:,.1f} t   "
                     f"Fire: {tr['fire_kg'] / 1000:,.1f} t")
        _header(wsr, 4, ["Çap", "Sipariş (kg)", "Sipariş (ton)", "Metraj (kg)", "Oranla (kg)", "Fire (kg)",
                         "Uzunluk (m)", "Nerede", "Kaynak"])
        for r in rebar["rows"]:
            wsr.append([f"Ø{r['dia_mm']}", round(r["order_kg"], 1), round(r["order_kg"] / 1000, 3), round(r["metraj_kg"], 1),
                        round(r["ratio_kg"], 1) or None, round(r["fire_kg"], 1) or None, round(r["length_m"], 1) or None,
                        "; ".join(f"{k} {v / 1000:,.1f} t" for k, v in sorted(r["targets"].items(), key=lambda kv: -kv[1])),
                        "; ".join(f"{k} {v / 1000:,.1f} t" for k, v in r["sources"].items())])
        u = rebar.get("unsized")
        if u:
            wsr.append(["çapsız", round(u["metraj_kg"] + u["ratio_kg"] + u["fire_kg"], 1), None, round(u["metraj_kg"], 1),
                        round(u["ratio_kg"], 1) or None, round(u["fire_kg"], 1) or None, None,
                        "çizimde donatı yazısı yok, çapa bölünemedi", "; ".join(u["groups"])])
        wsr.append(["TOPLAM", round(tr["order_kg"], 1), round(tr["order_kg"] / 1000, 3), round(tr["metraj_kg"], 1),
                    round(tr["ratio_kg"], 1) or None, round(tr["fire_kg"], 1) or None])
        for c in wsr[wsr.max_row]:
            c.font = BOLD
        wsr.append([])
        _header(wsr, wsr.max_row + 1, ["Demir işçiliği", "Saat"])
        for l in rebar.get("labour", []):
            wsr.append([l["label"], round(l["hours"], 1)])
        wsr.append(["TOPLAM", round(rebar.get("labour_hours", 0), 1)])
        for c in wsr[wsr.max_row]:
            c.font = BOLD
        if rebar.get("extras"):
            wsr.append([])
            _header(wsr, wsr.max_row + 1, ["Demire bağlı sarf", "Miktar", "Birim"])
            for e in rebar["extras"]:
                wsr.append([e["label"], e["quantity"], e["unit"]])
        _autosize(wsr)

    # ---- Kat (çizim) bazında ----
    if summary.get("by_drawing"):
        wsk = wb.create_sheet("Kat Bazında")
        etypes = ["foundation", "column", "shear_wall", "beam", "slab"]
        hdr = ["Plan / pafta", "Kot"]
        for et in etypes:
            hdr += [f"{ALL_ELEMENT_TYPES.get(et, et)} beton (m³)", f"{ALL_ELEMENT_TYPES.get(et, et)} kalıp (m²)"]
        hdr += ["Beton toplam (m³)", "Kalıp toplam (m²)", "Demir (oranla, kg)", "Demir tablo (kg)", "Tablo hedefi", "Çap bazında (kg)"]
        _header(wsk, 1, hdr)
        for r in summary["by_drawing"]:
            row = [r["drawing"], r.get("kot") or ""]
            for et in etypes:
                g = r["groups"].get(et, {})
                row += [g.get("concrete_m3", 0.0) or None, g.get("formwork_m2", 0.0) or None]
            row += [r["concrete_m3"], r["formwork_m2"], r["rebar_kg"] or None, r.get("rebar_table_kg"),
                    ALL_ELEMENT_TYPES.get(r.get("rebar_target", ""), r.get("rebar_target", "")) if r.get("rebar_table_kg") else "",
                    "; ".join(f"Ø{d}: {kg:,.0f}" for d, kg in (r.get("rebar_by_dia") or {}).items())]
            wsk.append(row)
        if summary.get("rebar_by_dia"):
            wsk.append([])
            _header(wsk, wsk.max_row + 1, ["Çap", "Toplam boy (m)", "Ağırlık (kg)", "Dağılım"])
            for d in summary["rebar_by_dia"]:
                wsk.append([f"Ø{d['dia_mm']}", d["length_m"], d["weight_kg"],
                            "; ".join(f"{ALL_ELEMENT_TYPES.get(k, k)} {v:,.0f}" for k, v in d["targets"].items())])
        _autosize(wsk)

    # ---- Eleman Listesi (statik) ----
    if lines:
        ws2 = wb.create_sheet("Eleman Metrajı")
        _header(ws2, 1, ["Çizim", "Tip", "Ad", "Katman", "b (m)", "h (m)", "Kalınlık (m)", "Alan (m²)",
                         "Uzunluk (m)", "Adet", "Kat Çarpanı", "Beton (m³)", "Kalıp (m²)", "Demir (kg)",
                         "Toplam Beton (m³)", "Toplam Kalıp (m²)", "Toplam Demir (kg)", "Notlar"])
        for ln in lines:
            info = element_info.get(ln.element_id, {})
            label = ALL_ELEMENT_TYPES.get(ln.etype, ln.etype)
            if ln.subtype:
                label = SUBTYPE_LABELS.get(ln.subtype, label)
            ws2.append([
                info.get("drawing", ""), label, ln.name or "", info.get("layer", ""),
                info.get("b"), info.get("h"), info.get("thickness"), info.get("area"), info.get("length"),
                ln.count, ln.multiplier, round(ln.concrete_m3, 4), round(ln.formwork_m2, 4), round(ln.rebar_kg, 2),
                round(ln.total_concrete, 4), round(ln.total_formwork, 4), round(ln.total_rebar, 2),
                "; ".join(ln.notes + list(info.get("warnings", []))),
            ])
        _autosize(ws2)

    # ---- Birim fiyatlar: malzeme ve işçilik ayrı sayfalarda ----
    # ürün bazında: C30/37 beton, Ø12 demir, Ytong 20 cm… (aynı ürünü kullanan kalemler tek satırda)
    wsm = wb.create_sheet("Malzeme Fiyatları")
    _header(wsm, 1, ["Ürün", "Birim", "Miktar", "Marka", "Birim Fiyat (₺)", "Tutar (₺)", "Kalem Sayısı", "Kullanan Kalemler"])
    used: dict[str, list[str]] = {}
    for l in cost["lines"]:
        if l.get("material_key"):
            used.setdefault(l["material_key"], []).append(l["group_label"])
    for m in cost.get("by_material", []):
        names = used.get(m["key"], [])
        wsm.append([m["name"], m["unit"], m["quantity"], m.get("brand", ""), m["unit_price"], m["total"], m["lines"],
                    "; ".join(names[:12]) + (f" (+{len(names) - 12})" if len(names) > 12 else "")])
    wsm.append([])
    wsm.cell(row=wsm.max_row + 1, column=5, value="Malzeme ara toplam").font = BOLD
    wsm.cell(row=wsm.max_row, column=6, value=cost.get("material_subtotal", 0.0)).font = BOLD
    for row in wsm.iter_rows(min_row=2, min_col=5, max_col=6):
        for c in row:
            if isinstance(c.value, (int, float)):
                c.number_format = MONEY
    _autosize(wsm)

    # kalem bazında malzeme dökümü (hangi kalem hangi üründen fiyatlandı)
    wsd = wb.create_sheet("Malzeme Dökümü")
    _header(wsd, 1, ["İş Grubu", "Poz", "Tür", "Kalem", "Birim", "Miktar", "Ürün", "Marka", "Birim Fiyat (₺)", "Tutar (₺)", "Kaynak"])
    for l in cost["lines"]:
        wsd.append([l.get("work_group_label", ""), l.get("poz", ""), l["kind_label"], l["group_label"], l["unit"], l["quantity"],
                    l.get("material_name", ""), l.get("brand", ""), l["unit_price"], l.get("material_total", 0.0), l["price_source"]])
    for row in wsd.iter_rows(min_row=2, min_col=9, max_col=10):
        for c in row:
            if isinstance(c.value, (int, float)):
                c.number_format = MONEY
    _autosize(wsd)

    wsl = wb.create_sheet("İşçilik Fiyatları")
    _header(wsl, 1, ["İş Grubu", "Poz", "Tür", "Kalem", "Birim", "Miktar", "İşçilik (₺/birim)", "İşçilik Tutarı (₺)",
                     "Adam-saat/birim", "Ekip", "Adam-saat", "Süre (gün)", "Kaynak"])
    for l in cost["lines"]:
        wsl.append([l.get("work_group_label", ""), l.get("poz", ""), l["kind_label"], l["group_label"], l["unit"], l["quantity"],
                    l.get("labor_price", 0.0), l.get("labor_total", 0.0), l.get("hours_per_unit", 0.0), l.get("crew_size", 1.0),
                    l.get("hours", 0.0), l.get("days", 0.0), l.get("labor_source", "")])
    wsl.append([])
    wsl.cell(row=wsl.max_row + 1, column=7, value="İşçilik ara toplam").font = BOLD
    wsl.cell(row=wsl.max_row, column=8, value=cost.get("labor_subtotal", 0.0)).font = BOLD
    for row in wsl.iter_rows(min_row=2, min_col=7, max_col=8):
        for c in row:
            if isinstance(c.value, (int, float)):
                c.number_format = MONEY
    _autosize(wsl)

    # ---- Maliyet ve Süre ----
    ws3 = wb.create_sheet("Maliyet")
    _header(ws3, 1, ["Disiplin", "Tür", "Kalem", "Ürün", "Birim", "Miktar", "Malzeme (₺/birim)", "İşçilik (₺/birim)",
                     "Malzeme Tutarı (₺)", "İşçilik Tutarı (₺)", "Toplam (₺)", "Adam-saat/birim", "Ekip", "Süre (gün)",
                     "Fiyat Kaynağı"])
    for l in cost["lines"]:
        ws3.append([l["discipline_label"], l["kind_label"], l["group_label"], l.get("material_name", ""), l["unit"], l["quantity"],
                    l["unit_price"], l.get("labor_price", 0.0), l.get("material_total", 0.0), l.get("labor_total", 0.0),
                    l["total"], l.get("hours_per_unit", 0.0), l.get("crew_size", 1.0), l.get("days", 0.0),
                    l["price_source"]])
    r = ws3.max_row + 2
    rows = [("Malzeme ara toplam", cost.get("material_subtotal", 0.0)), ("İşçilik ara toplam", cost.get("labor_subtotal", 0.0)),
            ("Ara Toplam", cost["subtotal"])]
    if cost.get("vat_rate"):
        rows.append((f"KDV (%{cost['vat_rate']*100:.0f})", cost["vat"]))
    rows.append(("GENEL TOPLAM", cost["grand_total"]))
    for label, val in rows:
        ws3.cell(row=r, column=10, value=label).font = BOLD
        c = ws3.cell(row=r, column=11, value=val)
        c.font = BOLD
        c.number_format = MONEY
        r += 1
    dur = cost.get("duration", {})
    r += 1
    ws3.cell(row=r, column=10, value="Toplam adam-saat").font = BOLD
    ws3.cell(row=r, column=11, value=dur.get("total_hours", 0.0))
    ws3.cell(row=r + 1, column=10, value="Süre, işler ardışık (gün)").font = BOLD
    ws3.cell(row=r + 1, column=11, value=dur.get("sequential_days", 0.0))
    ws3.cell(row=r + 2, column=10, value="Süre, disiplinler paralel (gün)").font = BOLD
    ws3.cell(row=r + 2, column=11, value=dur.get("parallel_days", 0.0))
    rr = r + 4
    _header(ws3, rr, ["Disiplin", "Malzeme (₺)", "İşçilik (₺)", "Toplam (₺)", "Adam-saat", "Süre (gün)"])
    for d in cost.get("by_discipline", []):
        rr += 1
        for ci, v in enumerate([d["label"], d["material"], d["labor"], d["total"], d["hours"], d["days"]], start=1):
            ws3.cell(row=rr, column=ci, value=v)
    for row in ws3.iter_rows(min_row=2, min_col=7, max_col=11):
        for c in row:
            if isinstance(c.value, (int, float)):
                c.number_format = MONEY
    _autosize(ws3)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
